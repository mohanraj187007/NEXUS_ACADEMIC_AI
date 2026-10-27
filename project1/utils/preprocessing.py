"""
Data Ingestion & Preprocessing Pipeline
Validates CSV and Excel files, handles dirty data, normalizes statuses, eliminates duplicates,
detects conflicting attendance entries, records row-level errors, and executes transaction-safe
database synchronization with complete data provenance and audit trails.
"""

import pandas as pd
import numpy as np
import io
import re
import uuid
import datetime
from typing import Dict, Any, Tuple, List, Optional
import sqlite3
from database.database import get_connection, log_audit_action, DB_PATH

COLUMN_ALIASES = {
    "student_id": ["student_id", "studentid", "roll_no", "reg_no", "rollno", "regno", "student_no", "id"],
    "student_name": ["student_name", "name", "full_name", "studentname", "candidate_name"],
    "register_number": ["register_number", "reg_number", "registration_no", "reg_no", "regno"],
    "department": ["department", "dept", "branch", "programme", "stream"],
    "semester": ["semester", "sem", "term"],
    "year": ["year", "batch_year", "yr", "academic_year"],
    "section": ["section", "sec", "division", "div", "batch"],
    "subject_id": ["subject_id", "subject", "subject_code", "course_code", "sub_code", "course_id"],
    "subject_name": ["subject_name", "course_name", "sub_name", "title"],
    "attendance_date": ["attendance_date", "date", "att_date", "session_date", "class_date"],
    "period_number": ["period_number", "period", "hour", "slot", "class_hour", "hr"],
    "scheduled_start_time": ["scheduled_start_time", "start_time", "class_start", "sched_start"],
    "scheduled_end_time": ["scheduled_end_time", "end_time", "class_end", "sched_end"],
    "attendance_status": ["attendance_status", "status", "presence", "att_status"],
    "check_in_time": ["check_in_time", "checkin", "in_time", "arrival_time", "punch_in"],
    "check_out_time": ["check_out_time", "checkout", "out_time", "departure_time", "punch_out"],
    "source_system": ["source_system", "source", "system", "vendor", "data_source"],
    "record_created_at": ["record_created_at", "created_at", "timestamp", "log_time"]
}

VALID_STATUSES = {
    "Present": ["present", "p", "1", "1.0", "yes", "true", "attended", "presente"],
    "Absent": ["absent", "a", "0", "0.0", "no", "false", "absente"],
    "Late": ["late", "l", "tardy", "delayed"],
    "Excused": ["excused", "ex", "medical", "duty_leave", "on_duty", "od"],
    "Approved Leave": ["approved leave", "leave", "approved_leave", "al", "on leave"],
    "Cancelled": ["cancelled", "class_cancelled", "c", "canceled"]
}


def normalize_column_headers(df: pd.DataFrame) -> pd.DataFrame:
    """Maps varying CSV/Excel header aliases to canonical database column names."""
    df_cols = {col: str(col).strip().lower().replace(" ", "_").replace("-", "_") for col in df.columns}
    df = df.rename(columns=df_cols)

    mapped_cols = {}
    for canonical_col, aliases in COLUMN_ALIASES.items():
        for col in df.columns:
            if col in aliases and canonical_col not in mapped_cols.values():
                mapped_cols[col] = canonical_col
                break

    return df.rename(columns=mapped_cols)


def normalize_status(val: Any) -> Tuple[Optional[str], bool]:
    """
    Normalizes status string to one of canonical statuses:
    'Present', 'Absent', 'Late', 'Excused', 'Approved Leave', 'Cancelled'.
    Returns (normalized_status, is_valid).
    """
    if pd.isna(val) or val is None:
        return ("Absent", True)  # Institutional default for blank if processed
    val_str = str(val).strip().lower()
    for canonical, aliases in VALID_STATUSES.items():
        if val_str in aliases:
            return (canonical, True)
    return (None, False)


def generate_csv_template() -> bytes:
    """Generates a downloadable CSV template with standard headers and sample demonstration rows."""
    template_cols = [
        "student_id", "student_name", "department", "semester", "year", "section",
        "subject_id", "subject_name", "attendance_date", "period_number",
        "scheduled_start_time", "scheduled_end_time", "attendance_status",
        "check_in_time", "check_out_time", "source_system"
    ]
    sample_rows = [
        ["STU_1001", "Aarav Gupta", "Computer Science", 5, 3, "A", "CS301", "Data Structures & Algorithms", "2024-03-01", 1, "09:00", "09:50", "Present", "08:58", "09:50", "FACULTY_ROSTER"],
        ["STU_1002", "Aditi Sundaram", "Computer Science", 5, 3, "A", "CS301", "Data Structures & Algorithms", "2024-03-01", 1, "09:00", "09:50", "Late", "09:12", "09:50", "BIOMETRIC_DOOR"],
        ["STU_1003", "Akash Sharma", "Information Technology", 5, 3, "B", "CS302", "Operating Systems", "2024-03-01", 2, "10:00", "10:50", "Absent", "", "", "FACULTY_ROSTER"],
        ["STU_1004", "Bhavya Reddy", "AI & Data Science", 3, 2, "A", "CS304", "Artificial Intelligence", "2024-03-01", 3, "11:10", "12:00", "Excused", "", "", "STUDENT_AFFAIRS_LEAVE"]
    ]
    df_temp = pd.DataFrame(sample_rows, columns=template_cols)
    return df_temp.to_csv(index=False).encode("utf-8")


def generate_excel_template() -> bytes:
    """Generates a downloadable Excel (.xlsx) template with styled headers."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Attendance_Import_Template"

    headers = [
        "student_id", "student_name", "department", "semester", "year", "section",
        "subject_id", "subject_name", "attendance_date", "period_number",
        "scheduled_start_time", "scheduled_end_time", "attendance_status",
        "check_in_time", "check_out_time", "source_system"
    ]
    ws.append(headers)

    sample_rows = [
        ["STU_1001", "Aarav Gupta", "Computer Science", 5, 3, "A", "CS301", "Data Structures & Algorithms", "2024-03-01", 1, "09:00", "09:50", "Present", "08:58", "09:50", "FACULTY_ROSTER"],
        ["STU_1002", "Aditi Sundaram", "Computer Science", 5, 3, "A", "CS301", "Data Structures & Algorithms", "2024-03-01", 1, "09:00", "09:50", "Late", "09:12", "09:50", "BIOMETRIC_DOOR"],
        ["STU_1003", "Akash Sharma", "Information Technology", 5, 3, "B", "CS302", "Operating Systems", "2024-03-01", 2, "10:00", "10:50", "Absent", "", "", "FACULTY_ROSTER"],
        ["STU_1004", "Bhavya Reddy", "AI & Data Science", 3, 2, "A", "CS304", "Artificial Intelligence", "2024-03-01", 3, "11:10", "12:00", "Excused", "", "", "STUDENT_AFFAIRS_LEAVE"]
    ]
    for row in sample_rows:
        ws.append(row)

    # Style header row
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def validate_and_process_file(
    file_bytes: bytes,
    file_name: str,
    conflict_strategy: str = "skip",  # "skip", "update_with_audit", or "reject_batch"
    authorized_user: str = "Admin",
    db_path: str = DB_PATH
) -> Tuple[bool, Dict[str, Any], pd.DataFrame]:
    """
    Parses CSV or Excel file, validates headers, runs row-level sanity tests,
    detects internal duplicates and database conflicts, executes a safe database transaction,
    and returns a structured audit summary and error register.
    """
    batch_id = f"BATCH_{uuid.uuid4().hex[:8].upper()}"
    now_iso = datetime.datetime.now().isoformat()

    # 1. Parse File Content
    try:
        if file_name.lower().endswith(".xlsx") or file_name.lower().endswith(".xls"):
            df_raw = pd.read_excel(io.BytesIO(file_bytes))
        else:
            try:
                df_raw = pd.read_csv(io.BytesIO(file_bytes), encoding="utf-8")
            except UnicodeDecodeError:
                df_raw = pd.read_csv(io.BytesIO(file_bytes), encoding="latin1")
    except Exception as e:
        return False, {
            "error": f"Failed to parse file: {str(e)}",
            "batch_id": batch_id,
            "total_raw": 0,
            "valid_records": 0,
            "invalid_records": 0,
            "row_errors": []
        }, pd.DataFrame()

    total_raw_rows = len(df_raw)
    if total_raw_rows == 0:
        return False, {
            "error": "The uploaded file is empty. Please provide attendance records.",
            "batch_id": batch_id,
            "total_raw": 0,
            "valid_records": 0,
            "invalid_records": 0,
            "row_errors": []
        }, pd.DataFrame()

    df = normalize_column_headers(df_raw.copy())

    # 2. Check for Essential Minimal Columns
    # Map synonyms
    if "attendance_date" in df.columns and "date" not in df.columns:
        df["date"] = df["attendance_date"]
    elif "date" in df.columns and "attendance_date" not in df.columns:
        df["attendance_date"] = df["date"]

    if "period_number" in df.columns and "hour" not in df.columns:
        df["hour"] = df["period_number"]
    elif "hour" in df.columns and "period_number" not in df.columns:
        df["period_number"] = df["hour"]

    if "attendance_status" in df.columns and "status" not in df.columns:
        df["status"] = df["attendance_status"]
    elif "status" in df.columns and "attendance_status" not in df.columns:
        df["attendance_status"] = df["status"]

    if "subject_id" in df.columns and "subject" not in df.columns:
        df["subject"] = df["subject_id"]
    elif "subject" in df.columns and "subject_id" not in df.columns:
        df["subject_id"] = df["subject"]

    required_cols = ["student_id", "attendance_date", "attendance_status"]
    missing = [c for c in required_cols if c not in df.columns]
    if missing:
        return False, {
            "error": f"Uploaded file is missing required columns: {', '.join(missing)}. Please verify template headers.",
            "batch_id": batch_id,
            "total_raw": total_raw_rows,
            "valid_records": 0,
            "invalid_records": total_raw_rows,
            "row_errors": [{"row": 1, "student_id": "N/A", "field": m, "error": "Missing column"} for m in missing]
        }, pd.DataFrame()

    # 3. Row-by-Row Deep Validation
    row_errors: List[Dict[str, Any]] = []
    valid_rows = []

    for idx, raw_row in df.iterrows():
        row_num = idx + 2  # 1-indexed Excel/CSV line number (header is row 1)
        s_id = str(raw_row.get("student_id", "")).strip()
        
        # Validate Student ID
        if not s_id or s_id.lower() in ["nan", "none", "null"]:
            row_errors.append({
                "row": row_num,
                "student_id": "EMPTY",
                "field": "student_id",
                "error": "Student ID cannot be empty or null.",
                "raw_val": str(raw_row.get("student_id", ""))
            })
            continue

        # Validate Date
        raw_date = raw_row.get("attendance_date")
        parsed_date = None
        try:
            parsed_date = pd.to_datetime(raw_date, errors="raise")
            date_str = parsed_date.strftime("%Y-%m-%d")
        except Exception:
            row_errors.append({
                "row": row_num,
                "student_id": s_id,
                "field": "attendance_date",
                "error": "Invalid date format. Expected YYYY-MM-DD.",
                "raw_val": str(raw_date)
            })
            continue

        # Validate Status
        raw_status = raw_row.get("attendance_status")
        norm_status, is_valid_status = normalize_status(raw_status)
        if not is_valid_status:
            row_errors.append({
                "row": row_num,
                "student_id": s_id,
                "field": "attendance_status",
                "error": f"Invalid attendance status '{raw_status}'. Must be Present, Absent, Late, Excused, Approved Leave, or Cancelled.",
                "raw_val": str(raw_status)
            })
            continue

        # Period / Hour
        raw_hour = raw_row.get("period_number", 1)
        try:
            hour_val = int(pd.to_numeric(raw_hour, errors="coerce")) if not pd.isna(raw_hour) else 1
            if hour_val < 1 or hour_val > 10:
                hour_val = max(1, min(10, hour_val))
        except Exception:
            hour_val = 1

        # Subject Code / ID
        raw_sub = str(raw_row.get("subject_id", raw_row.get("subject", "GEN101"))).strip()
        if not raw_sub or raw_sub.lower() in ["nan", "none"]:
            raw_sub = "GEN101"
        sub_id = f"SUB_{raw_sub}" if not raw_sub.startswith("SUB_") else raw_sub
        sub_code = raw_sub.replace("SUB_", "")
        sub_name = str(raw_row.get("subject_name", f"Course {sub_code}"))
        if sub_name.lower() in ["nan", "none"]:
            sub_name = f"Course {sub_code}"

        # Student Details Defaults
        stu_name = str(raw_row.get("student_name", raw_row.get("name", f"Student {s_id}"))).strip()
        if stu_name.lower() in ["nan", "none"]:
            stu_name = f"Student {s_id}"
        reg_no = str(raw_row.get("register_number", s_id)).strip()
        if reg_no.lower() in ["nan", "none"]:
            reg_no = s_id
        dept = str(raw_row.get("department", "General Engineering")).strip()
        if dept.lower() in ["nan", "none"]:
            dept = "General Engineering"
        
        sem = int(pd.to_numeric(raw_row.get("semester", 1), errors="coerce") or 1)
        yr = int(pd.to_numeric(raw_row.get("year", (sem + 1) // 2), errors="coerce") or 1)
        sec = str(raw_row.get("section", "A")).strip()
        if sec.lower() in ["nan", "none"]:
            sec = "A"

        # Timestamps
        sched_start = str(raw_row.get("scheduled_start_time", "")).strip() if not pd.isna(raw_row.get("scheduled_start_time")) else ""
        sched_end = str(raw_row.get("scheduled_end_time", "")).strip() if not pd.isna(raw_row.get("scheduled_end_time")) else ""
        check_in = str(raw_row.get("check_in_time", "")).strip() if not pd.isna(raw_row.get("check_in_time")) else ""
        check_out = str(raw_row.get("check_out_time", "")).strip() if not pd.isna(raw_row.get("check_out_time")) else ""
        source_sys = str(raw_row.get("source_system", "FILE_IMPORT")).strip() if not pd.isna(raw_row.get("source_system")) else "FILE_IMPORT"

        valid_rows.append({
            "student_id": s_id,
            "register_number": reg_no,
            "student_name": stu_name,
            "department": dept,
            "semester": sem,
            "year": yr,
            "section": sec,
            "subject_id": sub_id,
            "subject_code": sub_code,
            "subject_name": sub_name,
            "date": date_str,
            "status": norm_status,
            "hour": hour_val,
            "scheduled_start_time": sched_start,
            "scheduled_end_time": sched_end,
            "check_in_time": check_in,
            "check_out_time": check_out,
            "source_system": source_sys,
            "row_index": row_num
        })

    valid_df = pd.DataFrame(valid_rows)

    if valid_df.empty:
        return False, {
            "error": "No valid attendance rows found in the submitted file. Review row errors below.",
            "batch_id": batch_id,
            "total_raw": total_raw_rows,
            "valid_records": 0,
            "invalid_records": total_raw_rows,
            "row_errors": row_errors
        }, pd.DataFrame()

    # 4. Detect and Eliminate Internal Duplicates within the file
    before_internal_dedup = len(valid_df)
    valid_df = valid_df.drop_duplicates(subset=["student_id", "date", "hour", "subject_id"], keep="last").reset_index(drop=True)
    internal_duplicates_dropped = before_internal_dedup - len(valid_df)

    # Calculate academic week
    min_date = pd.to_datetime(valid_df["date"]).min()
    valid_df["week"] = ((pd.to_datetime(valid_df["date"]) - min_date).dt.days // 7) + 1

    # 5. Connect and Execute Safe Transactional Sync
    conn = get_connection(db_path)
    cursor = conn.cursor()

    inserted_count = 0
    updated_count = 0
    skipped_conflicts = 0

    try:
        conn.execute("BEGIN IMMEDIATE;")

        # A. Upsert Students
        unique_students = valid_df[["student_id", "register_number", "student_name", "department", "semester", "year", "section"]].drop_duplicates(subset=["student_id"])
        for _, s in unique_students.iterrows():
            cursor.execute("""
                INSERT INTO students (student_id, register_number, name, department, semester, year, section, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(student_id) DO UPDATE SET
                    register_number = excluded.register_number,
                    name = excluded.name,
                    department = excluded.department,
                    semester = excluded.semester,
                    year = excluded.year,
                    section = excluded.section;
            """, (s["student_id"], s["register_number"], s["student_name"], s["department"], int(s["semester"]), int(s["year"]), s["section"], now_iso))

        # B. Upsert Subjects
        unique_subjects = valid_df[["subject_id", "subject_code", "subject_name", "department", "semester"]].drop_duplicates(subset=["subject_id"])
        for _, sub in unique_subjects.iterrows():
            cursor.execute("""
                INSERT INTO subjects (subject_id, subject_code, subject_name, department, semester, faculty_name)
                VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(subject_id) DO UPDATE SET
                    subject_name = excluded.subject_name;
            """, (sub["subject_id"], sub["subject_code"], sub["subject_name"], sub["department"], int(sub["semester"]), "Faculty Assigned"))

        # C. Query Existing Records to Detect Database Conflicts
        # Check against existing attendance
        for _, row in valid_df.iterrows():
            cursor.execute("""
                SELECT attendance_id, status FROM attendance 
                WHERE student_id = ? AND date = ? AND hour = ? AND subject_id = ?
            """, (row["student_id"], row["date"], int(row["hour"]), row["subject_id"]))
            existing = cursor.fetchone()

            if existing is None:
                # Insert brand new record
                cursor.execute("""
                    INSERT INTO attendance (
                        student_id, subject_id, date, status, hour, week,
                        scheduled_start_time, scheduled_end_time, check_in_time, check_out_time,
                        source_system, record_created_at, updated_at, is_manually_corrected
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0);
                """, (
                    row["student_id"], row["subject_id"], row["date"], row["status"],
                    int(row["hour"]), int(row["week"]), row["scheduled_start_time"],
                    row["scheduled_end_time"], row["check_in_time"], row["check_out_time"],
                    row["source_system"], now_iso, now_iso
                ))
                inserted_count += 1
            else:
                existing_id, existing_status = existing
                if existing_status == row["status"]:
                    # Exact duplicate in database; skip
                    skipped_conflicts += 1
                else:
                    # Conflicting status!
                    if conflict_strategy == "reject_batch":
                        conn.rollback()
                        conn.close()
                        return False, {
                            "error": f"Conflicting attendance entry detected for student {row['student_id']} on {row['date']} period {row['hour']} (Existing: '{existing_status}', Incoming: '{row['status']}'). Batch transaction rejected.",
                            "batch_id": batch_id,
                            "total_raw": total_raw_rows,
                            "valid_records": len(valid_df),
                            "invalid_records": len(row_errors),
                            "row_errors": row_errors
                        }, pd.DataFrame()
                    elif conflict_strategy == "update_with_audit":
                        # Update record with audit trail
                        cursor.execute("""
                            UPDATE attendance SET 
                                status = ?, scheduled_start_time = ?, scheduled_end_time = ?,
                                check_in_time = ?, check_out_time = ?, source_system = ?,
                                updated_at = ?
                            WHERE attendance_id = ?;
                        """, (
                            row["status"], row["scheduled_start_time"], row["scheduled_end_time"],
                            row["check_in_time"], row["check_out_time"], row["source_system"],
                            now_iso, existing_id
                        ))
                        # Record in audit log
                        cursor.execute("""
                            INSERT INTO audit_logs (action_type, table_affected, record_id, old_value, new_value, performed_by, timestamp, details)
                            VALUES (?, ?, ?, ?, ?, ?, ?, ?);
                        """, (
                            "IMPORT_CONFLICT_UPDATE", "attendance", str(existing_id),
                            existing_status, row["status"], authorized_user, now_iso,
                            f"Batch {batch_id}: Updated via {file_name} for student {row['student_id']}"
                        ))
                        updated_count += 1
                    else:
                        # Skip conflicting record to preserve existing
                        skipped_conflicts += 1

        # Record overall batch ingestion in audit logs
        cursor.execute("""
            INSERT INTO audit_logs (action_type, table_affected, record_id, old_value, new_value, performed_by, timestamp, details)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?);
        """, (
            "FILE_INGESTION_BATCH", "attendance", batch_id, None,
            f"Inserted: {inserted_count}, Updated: {updated_count}, Skipped: {skipped_conflicts}",
            authorized_user, now_iso,
            f"File: {file_name}, Strategy: {conflict_strategy}, Raw rows: {total_raw_rows}"
        ))

        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        return False, {
            "error": f"Database transaction failed: {str(e)}",
            "batch_id": batch_id,
            "total_raw": total_raw_rows,
            "valid_records": len(valid_df),
            "invalid_records": len(row_errors),
            "row_errors": row_errors
        }, pd.DataFrame()
    finally:
        conn.close()

    summary = {
        "batch_id": batch_id,
        "file_name": file_name,
        "total_raw": total_raw_rows,
        "valid_records": len(valid_df),
        "inserted_records": inserted_count,
        "updated_records": updated_count,
        "skipped_conflicts": skipped_conflicts,
        "internal_duplicates_dropped": internal_duplicates_dropped,
        "invalid_records": len(row_errors),
        "row_errors": row_errors,
        "unique_students_imported": valid_df["student_id"].nunique(),
        "date_span": f"{valid_df['date'].min()} to {valid_df['date'].max()}"
    }

    return True, summary, valid_df


# Backward-compatible wrapper for existing test_suite and legacy callers
def validate_and_process_csv(df_raw: pd.DataFrame, db_path: str = DB_PATH) -> Tuple[bool, Dict[str, Any], pd.DataFrame]:
    """Backward-compatible wrapper function for legacy test calls."""
    csv_bytes = df_raw.to_csv(index=False).encode("utf-8")
    return validate_and_process_file(
        file_bytes=csv_bytes,
        file_name="legacy_upload.csv",
        conflict_strategy="skip",
        authorized_user="System",
        db_path=db_path
    )
